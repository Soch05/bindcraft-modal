#!/usr/bin/env python3
"""Rassemble les CSV d'une campagne BindCraft 2.0 en un classeur Excel par campagne.

openpyxl est une dépendance jetable, pas une étape du pipeline — même convention que
`build_workbook.py`, donc pas déclarée dans `pyproject.toml` :

    uv run --with openpyxl python csv_to_workbook.py out/egfr-dIII-prod02
    uv run --with openpyxl python csv_to_workbook.py out/*                 # toutes
    uv run --with openpyxl python csv_to_workbook.py out/egfr-dIII-prod02 --losses

Une feuille par table de campagne : `1_Trajectories`, `2_Refolded`, `3_Ranked`, `summary`.
Les ~90 `*_losses.csv` par trajectoire sont des traces de gradient ; elles sont **exclues par
défaut** parce que 90 feuilles de courbes ne se lisent pas. `--losses` les ajoute.

Les sorties vont dans `out/`, qui est gitignored : ce sont des rendus dérivés des CSV, et les
CSV restent la source. Même raisonnement que pour `data/*.xlsx`.
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1C5D82")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
MONO_FONT = Font(name="Menlo", size=9)

# Colonnes qu'on veut voir en premier, quand elles existent. Le reste suit dans l'ordre du
# CSV. `rank` et `design` d'abord parce que c'est par là qu'on lit une table de résultats.
PREFERRED_FIRST = (
    "rank",
    "trajectory",
    "design",
    "hash",
    "length",
    "terminated",
    "i_pTM",
    "i_pAE",
    "pLDDT",
    "pTM",
)

# Au-delà, une cellule devient illisible et Excel plafonne de toute façon à 32767.
MAX_CELL = 2000
MAX_WIDTH = 60


def sheet_name(raw: str) -> str:
    """Excel : 31 caractères maximum, et `[]:*?/\\` interdits."""
    clean = re.sub(r"[\[\]:*?/\\!]", "", raw).strip() or "feuille"
    return clean[:31]


def typed(value: str):
    """Rend les nombres numériques pour qu'Excel puisse trier et filtrer dessus."""
    text = value.strip()
    if not text or text.lower() in {"nan", "none", "inf", "-inf", "infinity"}:
        return None
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        return text[:MAX_CELL]


def add_sheet(book: Workbook, title: str, rows: list[dict[str, str]]) -> None:
    sheet = book.create_sheet(sheet_name(title))
    if not rows:
        sheet["A1"] = "(table vide)"
        return

    columns = list(rows[0])
    ordered = [c for c in PREFERRED_FIRST if c in columns]
    ordered += [c for c in columns if c not in ordered]

    sheet.append(ordered)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

    for row in rows:
        sheet.append([typed(row.get(column, "")) for column in ordered])

    # Les séquences et les listes de résidus sont illisibles dans une police
    # proportionnelle : on les passe en monospace.
    for index, column in enumerate(ordered, start=1):
        if any(key in column for key in ("Sequence", "Residues", "Timing", "hash")):
            for cell in sheet[get_column_letter(index)][1:]:
                cell.font = MONO_FONT

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for index, column in enumerate(ordered, start=1):
        widest = max(
            [len(column)] + [len(str(row.get(column, "") or "")) for row in rows[:400]]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(widest + 2, MAX_WIDTH)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def campaign_tables(root: Path, losses: bool) -> list[tuple[str, Path]]:
    """Les tables de campagne, dans l'ordre du pipeline, puis les traces si demandées."""
    tables: list[tuple[str, Path]] = []
    for stage in ("1_Trajectories", "2_Refolded", "3_Ranked"):
        for path in sorted((root / stage).glob("!_*.csv")):
            tables.append((stage, path))
    if (root / "summary.csv").is_file():
        tables.append(("summary", root / "summary.csv"))
    # Les tables de BindCraft 1, pour les runs historiques comme `smoke-G317`.
    for path in sorted(root.glob("*.csv")):
        if path.name != "summary.csv":
            tables.append((path.stem, path))
    if losses:
        for path in sorted(root.glob("1_Trajectories/*/*_losses.csv")):
            tables.append((path.parent.name[-28:], path))
    return tables


def selection_sheet(root: Path) -> list[dict[str, str]]:
    """Joint `3_Ranked`, les mutants acides et la conservation souris en une seule feuille.

    C'est la feuille de travail pour choisir ce qui part dans le CSV de soumission : la
    séquence d'origine et la séquence mutée côte à côte, avec les métriques et la
    conservation en regard.

    ⚠️ Les métriques décrivent la séquence **d'origine**. Celles d'un mutant sont inconnues
    et ne peuvent pas être obtenues avec `bindcraft score`, qui réaffiche le tampon du
    fichier au lieu de recalculer. Un mutant est une hypothèse à faire scorer ailleurs.
    """
    ranked = read_csv(root / "3_Ranked" / "!_Ranked.csv")
    if not ranked:
        return []
    mutants = {r["design"]: r for r in read_csv(Path("out/mutants_acide.csv"))}
    conservation = {
        r["pdb_resnum"]: r for r in read_csv(Path("data/egfr_residues.csv"))
    }

    rows = []
    for row in ranked:
        residues = re.findall(r"[A-Z](\d+)", row.get("Interface_Target_Residues", ""))
        identical = sum(
            1 for n in residues if conservation.get(n, {}).get("status") == "identical"
        )
        divergent = [
            f'{conservation[n]["aa_human"]}{n}>{conservation[n]["aa_mouse"]}'
            for n in residues
            if conservation.get(n, {}).get("status") == "different"
        ]
        mutant = mutants.get(f'{row["design"]}', {}) or mutants.get(
            f'{row["design"]}_seq0', {}
        )
        rows.append(
            {
                "design": row["design"],
                "hash_squelette": row["hash"],
                "length": row["length"],
                # objectif n°1
                "pH_etabli": mutant.get("pH_deja_etabli", ""),
                "pont_salin_A": mutant.get("pont_salin_existant_A", ""),
                "mutation_proposee": mutant.get("mutation", ""),
                "cout_substitution": mutant.get("cout_substitution", ""),
                "CB_vers_H409_A": mutant.get("CB_vers_H409_A", ""),
                # objectif n°2
                "souris_identiques": f"{identical}/{len(residues)}",
                "souris_divergents": ",".join(divergent) or "0",
                # objectif n°3 — sur la sequence d'ORIGINE uniquement
                "i_pTM": row.get("i_pTM", ""),
                "i_pAE": row.get("i_pAE", ""),
                "i_pDAE": row.get("i_pDAE", ""),
                "Interface_BuriedArea": row.get("Interface_BuriedArea", ""),
                "Interface_Residues": row.get("Interface_Residues", ""),
                "Hotspot_Contact_Fraction": row.get("Hotspot_Contact_Fraction", ""),
                "Unbound_Binder_pLDDT": row.get("Unbound_Binder_pLDDT", ""),
                "Surface_Hydrophobicity": row.get("Surface_Hydrophobicity", ""),
                "Binder_Net_Charge": row.get("Binder_Net_Charge", ""),
                # les deux sequences, cote a cote
                "Binder_Sequence": row.get("Binder_Sequence", ""),
                "Binder_Sequence_mutee": mutant.get("Binder_Sequence_mutee", ""),
            }
        )
    return rows


def build(root: Path, losses: bool) -> Path | None:
    tables = campaign_tables(root, losses)
    if not tables:
        print(f"  {root.name} : aucun CSV, ignoré")
        return None

    book = Workbook()
    book.remove(book.active)

    # La feuille de sélection d'abord : c'est celle qu'on ouvre pour travailler.
    selection = selection_sheet(root)
    if selection:
        add_sheet(book, "selection", selection)
        with_mutation = sum(1 for r in selection if r["Binder_Sequence_mutee"])
        print(f"    {'selection':<32} {len(selection):>5} lignes   "
              f"dont {with_mutation} avec sequence mutee")

    for title, path in tables:
        rows = read_csv(path)
        add_sheet(book, title, rows)
        print(f"    {title:<32} {len(rows):>5} lignes   {path.name}")

    destination = root / f"{root.name}.xlsx"
    book.save(destination)
    return destination


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("campaigns", type=Path, nargs="+", help="dossiers de campagne")
    parser.add_argument(
        "--losses",
        action="store_true",
        help="ajoute une feuille par trace de gradient (~90 feuilles, illisible)",
    )
    args = parser.parse_args()

    written = []
    for root in args.campaigns:
        if not root.is_dir():
            continue
        print(f"{root}")
        destination = build(root, args.losses)
        if destination:
            size_kb = destination.stat().st_size / 1024
            print(f"  -> {destination}  ({size_kb:.0f} Ko)")
            written.append(destination)
        print()

    print(f"{len(written)} classeur(s) écrit(s)")


if __name__ == "__main__":
    main()
