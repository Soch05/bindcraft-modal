#!/usr/bin/env python3
"""Construit le classeur Excel de travail pour choisir un site et ses hotspots.

Reprend les deux CSV du pipeline SANS les modifier, et y ajoute les colonnes
derivees qui manquaient pour trancher a la main :

  feuille `patches`   les 41 colonnes du CSV
                      + ancres_reelles / ancres_reelles_liste : les ancres acides
                        conservees dont le CARBOXYLATE passe le critere, et non les
                        ancres simplement annoncees
                      + ov_site_* : recouvrement de membres avec chacun des quatre
                        sites, pour voir l'appartenance au lieu de la deviner
                      + site : etiquette, vide si aucun recouvrement n'atteint
                        GROUP_LINK

  feuille `residus`   les 18 colonnes du CSV
                      + carbox_sasa / carbox_part / carbox_angle / ancre_reelle,
                        vides hors Asp et Glu

  feuille `ancres`    tous les Asp/Glu conserves de la chaine A avec leur verdict,
                      tries domaine III d'abord, puis ancres reelles, puis aire

  feuille `lisez-moi` legende des colonnes ajoutees, seuils et leur statut,
                      reserves

Toutes les valeurs sont des mesures importees, aucune n'est recalculee par Excel :
le classeur ne contient donc aucune formule. Les CSV restent la source suivie par
git ; ce classeur est un rendu, regenerable a tout moment.

Dependance jetable, non declaree dans pyproject.toml :
    uv run --with openpyxl python build_workbook.py
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import carboxylate_access as C
import egfr_epitope_map as E

OUT = E.DATA / "egfr_epitope_map.xlsx"
SITES = ("G317", "K375", "N449", "C502")
TEXT_COLUMNS = {"member_resnums_full", "ancres_reelles_liste", "site"}

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
ADDED_FILL = PatternFill("solid", fgColor="FFF2CC")  # colonnes derivees
HEADER_FONT = Font(name="Arial", size=10, bold=True)
BODY_FONT = Font(name="Arial", size=10)


def coerce(value: str, column: str):
    if column in TEXT_COLUMNS:
        return value
    if value in ("True", "False"):
        return value == "True"
    if value == "":
        return None
    try:
        return int(value)
    except ValueError:
        pass
    try:
        return float(value)
    except ValueError:
        return value


def measure_carboxylates() -> dict[int, dict]:
    """Mesures par residu acide, depuis la structure. Reutilise carboxylate_access."""
    entry = E.load_uniprot(E.HUMAN_AC)
    residues, _, _, _ = E.load_chain(entry["sequence"]["value"])
    out: dict[int, dict] = {}
    for res in residues:
        carb = C.carboxylate_sasa(res)
        if carb is None:
            continue
        total = float(sum(a.sasa for a in res))
        part = carb / total if total else None
        angle = C.outward_angle(res, residues)
        out[res.id[1]] = {
            "carbox_sasa": round(carb, 1),
            "carbox_part": round(part, 3) if part is not None else None,
            "carbox_angle": round(angle, 0) if angle is not None else None,
            "ancre_reelle": C.is_real_anchor(carb, part, angle),
        }
    return out


def site_unions(patches: list[dict]) -> dict[str, frozenset[int]]:
    """Union des membres de chaque site, via le groupement du pipeline."""
    by = {p["centre"]: p for p in patches}
    unions: dict[str, frozenset[int]] = {}
    for site in SITES:
        rep = by.get(site)
        if rep is None:
            continue
        grp = [rep] + [
            q
            for q in patches
            if q["centre"] != site
            and E._ov(q["member_set_full"], rep["member_set_full"]) >= E.GROUP_LINK
        ]
        unions[site] = frozenset().union(*[q["member_set_full"] for q in grp])
    return unions


def style(ws, nrow: int, header: list[str], added: set[str], widths: dict) -> None:
    for i, cell in enumerate(ws[1], start=1):
        cell.font = HEADER_FONT
        cell.fill = ADDED_FILL if header[i - 1] in added else HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for line in ws.iter_rows(min_row=2, max_row=nrow + 1, max_col=len(header)):
        for cell in line:
            cell.font = BODY_FONT
    for i, col in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 13)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{nrow + 1}"


def main() -> None:
    carbox = measure_carboxylates()

    # patches : CSV + colonnes derivees
    with (E.DATA / "egfr_patches.csv").open(newline="") as fh:
        reader = csv.DictReader(fh)
        pat_rows = list(reader)
        pat_header = list(reader.fieldnames or [])

    for r in pat_rows:
        r["member_set_full"] = frozenset(
            int(x) for x in re.findall(r"\d+", r["member_resnums_full"])
        )
    unions = site_unions(pat_rows)

    res_by_num = {
        int(r["pdb_resnum"]): r
        for r in csv.DictReader(open(E.DATA / "egfr_residues.csv"))
    }

    added_pat = ["ancres_reelles", "ancres_reelles_liste"]
    added_pat += [f"ov_site_{s}" for s in SITES] + ["site"]
    for r in pat_rows:
        members = sorted(r["member_set_full"])
        real = [
            n
            for n in members
            if res_by_num.get(n, {}).get("acidic") == "True"
            and res_by_num.get(n, {}).get("status") == "identical"
            and carbox.get(n, {}).get("ancre_reelle") is True
        ]
        r["ancres_reelles"] = str(len(real))
        r["ancres_reelles_liste"] = ";".join(
            f"{res_by_num[n]['aa_human']}{n}" for n in real
        )
        best, best_ov = "", 0.0
        for s, union in unions.items():
            ov = E._ov(r["member_set_full"], union)
            r[f"ov_site_{s}"] = f"{ov:.3f}"
            if ov > best_ov:
                best, best_ov = s, ov
        r["site"] = best if best_ov >= E.GROUP_LINK else ""

    # residus : CSV + mesures de carboxylate
    with (E.DATA / "egfr_residues.csv").open(newline="") as fh:
        reader = csv.DictReader(fh)
        res_rows = list(reader)
        res_header = list(reader.fieldnames or [])
    added_res = ["carbox_sasa", "carbox_part", "carbox_angle", "ancre_reelle"]
    for r in res_rows:
        m = carbox.get(int(r["pdb_resnum"]), {})
        for k in added_res:
            v = m.get(k)
            r[k] = "" if v is None else ("True" if v is True else ("False" if v is False else str(v)))

    wb = Workbook()
    wb.remove(wb.active)

    # --- lisez-moi ---
    ws = wb.create_sheet("lisez-moi")
    legend = [
        ("Classeur de travail — carte d'épitope EGFR, domaine III", ""),
        ("", ""),
        ("Source", "egfr_patches.csv et egfr_residues.csv, inchangés. Ce classeur est un rendu."),
        ("Colonnes sur fond jaune", "ajoutées ici, absentes des CSV. Dérivées, pas mesurées par Excel."),
        ("Formules", "aucune. Toutes les valeurs sont des mesures importées."),
        ("", ""),
        ("ancres_reelles", "ancres acides conservées dont le CARBOXYLATE passe le critère."),
        ("", "À ne pas confondre avec n_acidic_cons, qui compte les ancres annoncées."),
        ("ov_site_*", "fraction des membres du patch présents dans l'union de ce site."),
        ("site", "site dont le recouvrement est le plus fort, si >= GROUP_LINK. Vide sinon."),
        ("carbox_sasa", "SASA des deux oxygènes du carboxylate, en Å². Vide hors Asp/Glu."),
        ("carbox_part", "part du carboxylate dans la SASA du résidu."),
        ("carbox_angle", "angle CB->carboxylate contre le vecteur sortant local, en degrés."),
        ("ancre_reelle", "verdict composite sur les trois critères ci-dessous."),
        ("", ""),
        ("CRITÈRES D'ANCRE RÉELLE — POSÉS, NON CALIBRÉS", ""),
        (f"carbox_sasa >= {C.MIN_CARBOX_SASA:.0f} Å²",
         "section d'un cycle imidazole, ~4,5 x 4,0 Å. Argument géométrique."),
        (f"carbox_part >= {C.MIN_CARBOX_PART:.0%}",
         "distingue le groupe fonctionnel exposé de la seule tige."),
        (f"carbox_angle <= {C.MAX_OUTWARD_DEG:.0f}°",
         "sous 60° pointe vers le solvant ; au-delà de 90° longe la surface."),
        ("Calibration", "ÉCHOUÉE : l'empreinte du cétuximab ne contient qu'un Asp/Glu (E472)."),
        ("", ""),
        ("RÉSERVES", ""),
        ("Rotamère unique", "SASA sur un seul rotamère cristallographique à 3,20 Å."),
        ("Conformation repliée", "toute exposition est conditionnelle à l'état replié de 6ARU."),
        ("min_glyc", "mesuré au CB du séquon, pas à l'arbre glycanique. Sous-estime l'occlusion."),
        ("SASA apolaire", "borne la designabilité, pas l'affinité."),
        ("Plancher 400 Å²", "référence d'échelle, pas seuil de suffisance. Un Fab enfouit plus."),
        ("PATCH_RADIUS = 11 Å", "posé, non calibré. Capture ~40 % d'une empreinte réelle."),
        ("", ""),
        ("Journal complet", "NOTES.md — toutes les décisions et leurs motifs."),
        ("Colonne par colonne", "data/LECTURE.md"),
    ]
    for a, b in legend:
        ws.append([a, b])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 86
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for r in (1, 16, 22):
        ws.cell(row=r, column=1).font = Font(name="Arial", size=10, bold=True)

    # --- patches ---
    head = pat_header + added_pat
    ws = wb.create_sheet("patches")
    ws.append(head)
    for r in pat_rows:
        ws.append([coerce(str(r.get(c, "")), c) for c in head])
    style(ws, len(pat_rows), head, set(added_pat),
          {"centre": 9, "member_resnums_full": 30, "ancres_reelles_liste": 18, "site": 9})

    # --- residus ---
    head = res_header + added_res
    ws = wb.create_sheet("residus")
    ws.append(head)
    for r in res_rows:
        ws.append([coerce(str(r.get(c, "")), c) for c in head])
    style(ws, len(res_rows), head, set(added_res), {"status": 11})

    # --- ancres ---
    ws = wb.create_sheet("ancres")
    head = ["residu", "pdb", "uniprot", "carbox_sasa", "carbox_part", "carbox_angle",
            "ancre_reelle", "in_domain3", "dist_glycan", "dist_fab", "sites"]
    ws.append(head)
    acid = [
        r for r in res_rows
        if r["acidic"] == "True" and r["status"] == "identical" and r["carbox_sasa"] != ""
    ]
    site_of = {}
    for r in pat_rows:
        for n in r["member_set_full"]:
            if r["site"]:
                site_of.setdefault(n, set()).add(r["site"])
    # domaine III d'abord : c'est le seul perimetre de design. Puis verdict, puis aire.
    acid.sort(
        key=lambda r: (
            r["in_domain3"] != "True",
            r["ancre_reelle"] != "True",
            -float(r["carbox_sasa"]),
        )
    )
    for r in acid:
        n = int(r["pdb_resnum"])
        ws.append([
            f"{r['aa_human']}{n}", n, int(r["uniprot_pos"]),
            float(r["carbox_sasa"]), float(r["carbox_part"]), float(r["carbox_angle"]),
            r["ancre_reelle"] == "True", r["in_domain3"] == "True",
            float(r["dist_glycan"]), float(r["dist_fab"]),
            ";".join(sorted(site_of.get(n, []))),
        ])
    style(ws, len(acid), head, {"carbox_sasa", "carbox_part", "carbox_angle", "ancre_reelle"},
          {"residu": 10, "sites": 14})

    wb.save(OUT)
    print(f"ecrit {OUT}")
    print(f"  patches   {len(pat_rows)} lignes x {len(pat_header) + len(added_pat)} colonnes")
    print(f"  residus   {len(res_rows)} lignes x {len(res_header) + len(added_res)} colonnes")
    print(f"  ancres    {len(acid)} lignes")
    print(f"  ancres reelles : {sum(1 for r in acid if r['ancre_reelle'] == 'True')}/{len(acid)}")


if __name__ == "__main__":
    main()
